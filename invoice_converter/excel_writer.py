from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from .models import InvoiceRecord


HEADERS = [
    "序号",
    "发票类型",
    "发票代码",
    "发票号码",
    "开票日期",
    "购买方信息",
    "销售方信息",
    "不含税金额",
    "税额",
    "价税合计",
    "税率",
    "项目名称",
    "备注",
    "收款人",
    "复核人",
    "开票人",
    "来源文件",
    "页码",
    "识别方式",
    "OCR置信度",
    "识别状态",
    "异常说明",
]


def _safe_text(value: str) -> str:
    """Prevent OCR-controlled text from becoming an Excel formula."""
    if value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _party_info(name: str, tax_id: str) -> str:
    return f"名称：{name}\n纳税人识别号：{tax_id}"


def _row(record: InvoiceRecord, index: int) -> list:
    return [
        index,
        _safe_text(record.invoice_type),
        _safe_text(record.invoice_code),
        _safe_text(record.invoice_number),
        record.invoice_date,
        _safe_text(_party_info(record.buyer_name, record.buyer_tax_id)),
        _safe_text(_party_info(record.seller_name, record.seller_tax_id)),
        record.amount_without_tax,
        record.tax_amount,
        record.total_amount,
        _safe_text(record.tax_rate),
        _safe_text(record.item_name),
        _safe_text(record.remarks),
        _safe_text(record.payee),
        _safe_text(record.reviewer),
        _safe_text(record.issuer),
        _safe_text(record.source_file),
        record.page_number,
        _safe_text(record.recognition_method),
        record.ocr_confidence,
        _safe_text(record.status),
        _safe_text(record.issues),
    ]


def write_excel(records: list[InvoiceRecord], output_path: str | Path) -> Path:
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "发票明细"
    sheet.sheet_view.showGridLines = False
    sheet.append(HEADERS)
    for index, record in enumerate(records, start=1):
        sheet.append(_row(record, index))

    dark_blue = "1F4E78"
    light_blue = "D9EAF7"
    thin_gray = Side(style="thin", color="D9E2F3")
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor=dark_blue)
        cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=thin_gray)
    sheet.row_dimensions[1].height = 28

    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Microsoft YaHei", size=10, color="222222")
            cell.alignment = Alignment(vertical="center")
        row[4].number_format = "yyyy-mm-dd"
        for position in (7, 8, 9):
            row[position].number_format = '#,##0.00'
        row[5].alignment = Alignment(vertical="center", wrap_text=True)
        row[6].alignment = Alignment(vertical="center", wrap_text=True)
        row[19].number_format = "0.0%"
        sheet.row_dimensions[row[0].row].height = 38

    if records:
        table = Table(displayName="InvoiceDetails", ref=f"A1:V{len(records) + 1}")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        sheet.add_table(table)
        warning_fill = PatternFill("solid", fgColor="FFF2CC")
        error_fill = PatternFill("solid", fgColor="FCE4D6")
        sheet.conditional_formatting.add(
            f"U2:U{len(records) + 1}",
            FormulaRule(formula=['U2="需复核"'], fill=warning_fill),
        )
        sheet.conditional_formatting.add(
            f"U2:U{len(records) + 1}",
            FormulaRule(formula=['U2="失败"'], fill=error_fill),
        )

    widths = {
        1: 8, 2: 24, 3: 16, 4: 24, 5: 13, 6: 42, 7: 42, 8: 14, 9: 12,
        10: 14, 11: 14, 12: 32, 13: 30, 14: 12, 15: 12, 16: 12, 17: 42,
        18: 8, 19: 12, 20: 12, 21: 12, 22: 36,
    }
    for column, width in widths.items():
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:V{max(1, len(records) + 1)}"
    sheet.sheet_properties.tabColor = dark_blue
    sheet.print_title_rows = "1:1"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_margins.left = 0.25
    sheet.page_margins.right = 0.25
    workbook.save(output)
    return output.resolve()
