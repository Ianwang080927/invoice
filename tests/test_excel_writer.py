from datetime import date
from pathlib import Path
import tempfile
import unittest

from openpyxl import load_workbook

from invoice_converter.excel_writer import HEADERS, write_excel
from invoice_converter.models import InvoiceRecord


class ExcelWriterTests(unittest.TestCase):
    def test_write_excel(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            output = Path(temp_dir) / "result.xlsx"
            write_excel(
                [
                    InvoiceRecord(
                        invoice_number="12345678",
                        invoice_date=date(2026, 9, 18),
                        buyer_name="购买方公司",
                        buyer_tax_id="BUYER123456789",
                        seller_name="测试公司",
                        seller_tax_id="SELLER123456789",
                        total_amount=106.0,
                        source_file="sample.jpg",
                    )
                ],
                output,
            )
            workbook = load_workbook(output)
            sheet = workbook["发票明细"]
            self.assertEqual([cell.value for cell in sheet[1]], HEADERS)
            self.assertEqual(sheet["D2"].value, "12345678")
            self.assertEqual(sheet["F2"].value, "名称：购买方公司\n纳税人识别号：BUYER123456789")
            self.assertEqual(sheet["G2"].value, "名称：测试公司\n纳税人识别号：SELLER123456789")
            self.assertTrue(sheet["F2"].alignment.wrap_text)
            self.assertTrue(sheet["G2"].alignment.wrap_text)
            self.assertEqual(sheet["J2"].value, 106.0)
            self.assertEqual(sheet.freeze_panes, "A2")
            self.assertEqual(next(iter(sheet.tables.values())).ref, "A1:V2")

    def test_ocr_text_cannot_create_excel_formula(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            output = Path(temp_dir) / "safe.xlsx"
            write_excel([InvoiceRecord(item_name="=WEBSERVICE(\"bad\")")], output)
            sheet = load_workbook(output, data_only=False)["发票明细"]
            self.assertEqual(sheet["L2"].data_type, "s")
            self.assertTrue(sheet["L2"].value.startswith("'="))


if __name__ == "__main__":
    unittest.main()
